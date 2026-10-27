"""Reader checks for digital PDF, DOCX, scanned pages, images, and limits."""

from __future__ import annotations

import os
import tempfile
from io import BytesIO
from pathlib import Path

import pytest
from PIL import Image
from docx import Document

import report_documents
from report_documents import (
    DocumentLimitError,
    DocumentLimits,
    DocumentReadError,
    ExtractionMethod,
    ImageReference,
    ReadStatus,
    TesseractOcrAdapter,
    UnsupportedFormatError,
    read_report_document,
)


class FakeOcr:
    """Test double for the OCR adapter. It does not shell out to Tesseract."""

    def __init__(self, text: str = "Scanned rail fracture at KM 67.5.", *, fail: bool = False, ready: bool = True) -> None:
        self.text = text
        self.fail = fail
        self.ready = ready
        self.calls: list[str] = []

    def available(self) -> bool:
        return self.ready

    def read_image(self, image: ImageReference) -> str:
        self.calls.append(image.image_id)
        if self.fail:
            raise RuntimeError("ocr engine crashed")
        return self.text


def _stream(payload: bytes) -> bytes:
    return b"<< /Length %d >>\nstream\n" % len(payload) + payload + b"\nendstream"


def _build_pdf(objects: list[bytes]) -> bytes:
    header = b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n"
    chunks = [header]
    offsets: list[int] = []
    for index, obj in enumerate(objects, start=1):
        offsets.append(sum(len(chunk) for chunk in chunks))
        chunks.append(f"{index} 0 obj\n".encode("ascii") + obj + b"\nendobj\n")
    xref_at = sum(len(chunk) for chunk in chunks)
    xref = [f"xref\n0 {len(objects) + 1}\n".encode("ascii"), b"0000000000 65535 f \n"]
    for offset in offsets:
        xref.append(f"{offset:010d} 00000 n \n".encode("ascii"))
    trailer = (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n"
    ).encode("ascii")
    return b"".join(chunks + xref + [trailer])


def _text_stream(text: str) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    return _stream(f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("ascii"))


def _font() -> bytes:
    return b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"


def _catalog() -> bytes:
    return b"<< /Type /Catalog /Pages 2 0 R >>"


def _text_pdf(pages: list[str]) -> bytes:
    font_id = 3 + (2 * len(pages))
    objects: list[bytes] = [
        _catalog(),
        f"<< /Type /Pages /Count {len(pages)} /Kids [{' '.join(f'{3 + (2 * index)} 0 R' for index in range(len(pages)))}] >>".encode(
            "ascii"
        ),
    ]
    for index, text in enumerate(pages):
        page_id = 3 + (2 * index)
        content_id = page_id + 1
        objects.append(
            (
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_id} 0 R "
                f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>"
            ).encode("ascii")
        )
        objects.append(_text_stream(text))
    objects.append(_font())
    return _build_pdf(objects)


def _jpeg_bytes() -> bytes:
    buffer = BytesIO()
    Image.new("RGB", (4, 3), (10, 20, 30)).save(buffer, format="JPEG", quality=90)
    return buffer.getvalue()


def _png_bytes() -> bytes:
    buffer = BytesIO()
    Image.new("RGB", (6, 4), (1, 2, 3)).save(buffer, format="PNG")
    return buffer.getvalue()


def _image_object(jpeg: bytes) -> bytes:
    header = (
        f"<< /Type /XObject /Subtype /Image /Width 4 /Height 3 /ColorSpace /DeviceRGB "
        f"/BitsPerComponent 8 /Filter /DCTDecode /Length {len(jpeg)} >>\nstream\n"
    ).encode("ascii")
    return header + jpeg + b"\nendstream"


def _scanned_pdf(jpeg: bytes) -> bytes:
    return _build_pdf(
        [
            _catalog(),
            b"<< /Type /Pages /Count 1 /Kids [3 0 R] >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /XObject << /Im1 5 0 R >> >> >>",
            _stream(b"q 100 0 0 80 50 50 cm /Im1 Do Q"),
            _image_object(jpeg),
        ]
    )


def _text_and_image_pdf(text: str, jpeg: bytes) -> bytes:
    drawing = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET\nq 40 0 0 30 50 50 cm /Im1 Do Q".encode("ascii")
    return _build_pdf(
        [
            _catalog(),
            b"<< /Type /Pages /Count 1 /Kids [3 0 R] >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 6 0 R >> /XObject << /Im1 5 0 R >> >> >>",
            _stream(drawing),
            _image_object(jpeg),
            _font(),
        ]
    )


def _mixed_pdf(text: str, jpeg: bytes) -> bytes:
    return _build_pdf(
        [
            _catalog(),
            b"<< /Type /Pages /Count 2 /Kids [3 0 R 6 0 R] >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
            _text_stream(text),
            _font(),
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 7 0 R /Resources << /XObject << /Im1 8 0 R >> >> >>",
            _stream(b"q 100 0 0 80 50 50 cm /Im1 Do Q"),
            _image_object(jpeg),
        ]
    )


def _track_temp_dirs(monkeypatch: pytest.MonkeyPatch) -> list[str]:
    created: list[str] = []
    real = tempfile.TemporaryDirectory

    class TrackingTemporaryDirectory(real):  # type: ignore[type-arg, misc]
        def __init__(self, *args: object, **kwargs: object) -> None:
            super().__init__(*args, **kwargs)
            created.append(self.name)

    monkeypatch.setattr(report_documents.tempfile, "TemporaryDirectory", TrackingTemporaryDirectory)
    return created


def test_digital_pdf_extracts_selectable_text() -> None:
    sentence = "Rail inspection at KM 102.3 on the DN line."
    ocr = FakeOcr()
    bundle = read_report_document(_text_pdf([sentence]), "inspection.pdf", ocr=ocr)

    assert bundle.media_kind == "pdf"
    assert bundle.status is ReadStatus.COMPLETE
    assert bundle.warnings == ()
    assert len(bundle.sections) == 1
    section = bundle.sections[0]
    assert section.page_number == 1
    assert section.reference == "page:1"
    assert section.native_text == sentence
    assert section.ocr_text == ""
    assert section.blocks[0].method is ExtractionMethod.NATIVE_TEXT
    assert section.blocks[0].text == sentence
    assert ExtractionMethod.OCR not in section.methods
    assert ocr.calls == []


def test_multipage_pdf_keeps_page_numbers() -> None:
    first = "Page one records a weld failure at KM 16.8."
    second = "Page two records deep screening at KM 156.0."
    bundle = read_report_document(_text_pdf([first, second]), "two-pages.pdf", ocr=FakeOcr())

    assert [section.page_number for section in bundle.sections] == [1, 2]
    assert [section.reference for section in bundle.sections] == ["page:1", "page:2"]
    assert bundle.sections[0].native_text == first
    assert bundle.sections[1].native_text == second
    assert first not in bundle.sections[1].native_text
    assert second not in bundle.sections[0].native_text


def test_docx_preserves_paragraph_and_table_order(tmp_path: Path) -> None:
    path = tmp_path / "note.docx"
    document = Document()
    document.add_paragraph("Inspection note at KM 15.2.")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Asset"
    table.rows[0].cells[1].text = "Line"
    table.rows[1].cells[0].text = "Rail"
    table.rows[1].cells[1].text = "DN"
    document.add_paragraph("Closing note.")
    document.save(path)
    ocr = FakeOcr()

    bundle = read_report_document(path.read_bytes(), "note.docx", ocr=ocr)

    assert bundle.media_kind == "docx"
    assert bundle.status is ReadStatus.COMPLETE
    assert [section.reference for section in bundle.sections] == ["paragraph:1", "table:1", "paragraph:2"]
    assert all(section.page_number is None for section in bundle.sections)
    assert bundle.sections[0].native_text == "Inspection note at KM 15.2."
    assert bundle.sections[0].blocks[0].method is ExtractionMethod.NATIVE_TEXT
    assert bundle.sections[1].blocks[0].rows == (("Asset", "Line"), ("Rail", "DN"))
    assert bundle.sections[2].native_text == "Closing note."
    assert ocr.calls == []


def test_scanned_pdf_uses_ocr_adapter() -> None:
    jpeg = _jpeg_bytes()
    ocr = FakeOcr("Scanned rail fracture at KM 67.5.")
    bundle = read_report_document(_scanned_pdf(jpeg), "scan.pdf", ocr=ocr)

    section = bundle.sections[0]
    assert bundle.status is ReadStatus.COMPLETE
    assert section.page_number == 1
    assert section.native_text == ""
    assert section.ocr_text == "Scanned rail fracture at KM 67.5."
    assert section.blocks[0].method is ExtractionMethod.OCR
    assert section.blocks[0].reference == "page:1/image:1"
    assert ExtractionMethod.NATIVE_TEXT not in section.methods
    assert ExtractionMethod.IMAGE in section.methods
    assert ocr.calls == ["page:1/image:1"]


def test_ocr_unavailable_names_the_unread_page() -> None:
    jpeg = _jpeg_bytes()
    ocr = FakeOcr(ready=False)
    bundle = read_report_document(
        _mixed_pdf("Digital text on page one is long enough to keep.", jpeg),
        "mixed.pdf",
        ocr=ocr,
    )

    assert bundle.status is ReadStatus.PARTIAL
    assert bundle.sections[0].page_number == 1
    assert bundle.sections[0].native_text == "Digital text on page one is long enough to keep."
    assert bundle.sections[0].ocr_text == ""
    assert bundle.sections[1].page_number == 2
    assert bundle.sections[1].native_text == ""
    assert bundle.sections[1].ocr_text == ""
    assert bundle.sections[1].images[0].image_id == "page:2/image:1"
    unread = [warning for warning in bundle.warnings if warning.code == "OCR_UNAVAILABLE"]
    assert [warning.page_number for warning in unread] == [2]
    assert "Page 2 could not be read" in unread[0].message
    assert ocr.calls == []


def test_ocr_failure_names_the_unread_page() -> None:
    ocr = FakeOcr(fail=True)
    bundle = read_report_document(_scanned_pdf(_jpeg_bytes()), "scan.pdf", ocr=ocr)

    assert bundle.status is ReadStatus.PARTIAL
    assert bundle.sections[0].ocr_text == ""
    assert bundle.sections[0].images[0].data.startswith(b"\xff\xd8\xff")
    warning = bundle.warnings[0]
    assert warning.code == "OCR_FAILED"
    assert warning.page_number == 1
    assert "Page 1 could not be read" in warning.message
    assert "ocr engine crashed" in warning.message
    assert ocr.calls == ["page:1/image:1"]


def test_corrupt_pdf_raises_and_cleans_temp_files(monkeypatch: pytest.MonkeyPatch) -> None:
    created = _track_temp_dirs(monkeypatch)
    with pytest.raises(DocumentReadError) as caught:
        read_report_document(b"%PDF-1.4\nthis is not a pdf", "broken.pdf")
    assert type(caught.value) is DocumentReadError
    assert "broken.pdf" in str(caught.value)
    assert created
    assert all(not os.path.exists(path) for path in created)


def test_unsupported_format_is_rejected() -> None:
    with pytest.raises(UnsupportedFormatError):
        read_report_document(b"plain notes", "notes.txt")


def test_content_does_not_match_extension() -> None:
    with pytest.raises(UnsupportedFormatError):
        read_report_document(_png_bytes(), "report.pdf")


def test_embedded_image_keeps_a_stable_source_reference() -> None:
    jpeg = _jpeg_bytes()
    payload = _scanned_pdf(jpeg)
    ocr = FakeOcr("photo text")
    first = read_report_document(payload, "scan.pdf", ocr=ocr)
    second = read_report_document(payload, "scan.pdf", ocr=FakeOcr("photo text"))

    image = first.sections[0].images[0]
    assert image.image_id == "page:1/image:1"
    assert image.reference == "page:1/image:1"
    assert image.page_number == 1
    assert image.media_type == "image/jpeg"
    assert image.origin == "embedded_stream"
    assert image.data == jpeg
    assert image.sha256 == second.sections[0].images[0].sha256
    assert first.sections[0].images[0].image_id == second.sections[0].images[0].image_id


def test_docx_embedded_image_uses_paragraph_reference(tmp_path: Path) -> None:
    png = _png_bytes()
    image_path = tmp_path / "photo.png"
    image_path.write_bytes(png)
    document = Document()
    document.add_picture(str(image_path))
    target = tmp_path / "photo.docx"
    document.save(target)
    payload = target.read_bytes()

    bundle = read_report_document(payload, "photo.docx", ocr=FakeOcr("weld photo"))
    again = read_report_document(payload, "photo.docx", ocr=FakeOcr("weld photo"))
    image = bundle.sections[0].images[0]

    assert bundle.sections[0].reference == "paragraph:1"
    assert bundle.sections[0].page_number is None
    assert image.image_id == "paragraph:1/image:1"
    assert image.data == png
    assert image.origin == "embedded_part"
    assert image.image_id == again.sections[0].images[0].image_id
    assert image.sha256 == again.sections[0].images[0].sha256


def test_standalone_image_keeps_bytes_and_uses_ocr() -> None:
    png = _png_bytes()
    ocr = FakeOcr("uploaded sleeper photograph")
    bundle = read_report_document(png, "site.png", ocr=ocr)

    assert bundle.media_kind == "png"
    assert bundle.status is ReadStatus.COMPLETE
    assert bundle.sections[0].reference == "image:1"
    assert bundle.sections[0].page_number is None
    assert bundle.sections[0].native_text == ""
    assert bundle.sections[0].ocr_text == "uploaded sleeper photograph"
    assert bundle.sections[0].images[0].data == png
    assert bundle.sections[0].images[0].origin == "embedded_file"
    assert ocr.calls == ["image:1"]


def test_native_text_page_with_image_does_not_call_ocr() -> None:
    sentence = "Rail inspection at KM 102.3 on the DN line."
    ocr = FakeOcr()
    bundle = read_report_document(_text_and_image_pdf(sentence, _jpeg_bytes()), "logo.pdf", ocr=ocr)

    section = bundle.sections[0]
    assert section.native_text.rstrip("\n") == sentence
    assert section.ocr_text == ""
    assert section.images[0].image_id == "page:1/image:1"
    assert ExtractionMethod.OCR not in section.methods
    assert bundle.status is ReadStatus.COMPLETE
    assert ocr.calls == []


def test_short_native_text_without_an_image_is_partial() -> None:
    bundle = read_report_document(_text_pdf(["KM 15"]), "short.pdf", ocr=FakeOcr())

    assert bundle.status is ReadStatus.PARTIAL
    assert bundle.sections[0].native_text == "KM 15"
    warning = bundle.warnings[0]
    assert warning.code == "INSUFFICIENT_TEXT"
    assert warning.page_number == 1
    assert "Page 1 could not be read" in warning.message


def test_file_size_limit_rejects_before_parsing(monkeypatch: pytest.MonkeyPatch) -> None:
    created = _track_temp_dirs(monkeypatch)
    payload = b"%PDF-1.4\n" + (b"0" * 200)
    with pytest.raises(DocumentLimitError) as caught:
        read_report_document(payload, "big.pdf", limits=DocumentLimits(max_bytes=64))
    assert "64" in str(caught.value)
    assert created == []


def test_page_count_limit_cleans_temp_files(monkeypatch: pytest.MonkeyPatch) -> None:
    created = _track_temp_dirs(monkeypatch)
    payload = _text_pdf(["First page has enough text.", "Second page has enough text."])
    with pytest.raises(DocumentLimitError) as caught:
        read_report_document(payload, "long.pdf", limits=DocumentLimits(max_pdf_pages=1))
    assert "2 pages" in str(caught.value)
    assert "limit of 1 page" in str(caught.value)
    assert created
    assert all(not os.path.exists(path) for path in created)


def test_tesseract_adapter_reports_unavailable_without_program(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(report_documents.shutil, "which", lambda _name: None)
    assert TesseractOcrAdapter().available() is False


def test_success_path_cleans_temp_files(monkeypatch: pytest.MonkeyPatch) -> None:
    created = _track_temp_dirs(monkeypatch)
    bundle = read_report_document(
        _text_pdf(["Rail inspection at KM 102.3 on the DN line."]),
        "inspection.pdf",
        ocr=FakeOcr(),
    )
    assert bundle.status is ReadStatus.COMPLETE
    assert created
    assert all(not os.path.exists(path) for path in created)
