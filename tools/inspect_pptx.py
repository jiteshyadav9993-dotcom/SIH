"""
inspect_pptx.py
====================================================================
One-off diagnostic script (NOT part of the SIH-26027 DSS application).
Dumps the full slide/shape/text/font structure of a .pptx file to a
plain-text report so the structure can be reviewed before editing.
====================================================================
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu


def emu_to_in(value: int | None) -> str:
    if value is None:
        return "None"
    return f"{Emu(value).inches:.2f}in"


def describe_shape(shape, indent: str = "    ") -> list[str]:
    lines: list[str] = []
    lines.append(
        f"{indent}Shape id={shape.shape_id} name={shape.name!r} type={shape.shape_type} "
        f"pos=(L{emu_to_in(shape.left)},T{emu_to_in(shape.top)}) "
        f"size=(W{emu_to_in(shape.width)},H{emu_to_in(shape.height)})"
    )
    if shape.has_text_frame:
        for p_idx, para in enumerate(shape.text_frame.paragraphs):
            run_descs = []
            for r_idx, run in enumerate(para.runs):
                font = run.font
                color = None
                try:
                    if font.color and font.color.type is not None:
                        color = font.color.rgb
                except Exception:
                    color = "N/A"
                run_descs.append(
                    f"run{r_idx}[text={run.text!r} font={font.name!r} size={font.size} "
                    f"bold={font.bold} color={color}]"
                )
            if para.runs or para.text:
                lines.append(f"{indent}  para{p_idx} level={para.level} text={para.text!r}")
                for rd in run_descs:
                    lines.append(f"{indent}    {rd}")
    if shape.shape_type == 13:  # PICTURE
        lines.append(f"{indent}  [PICTURE]")
    if shape.has_table:
        table = shape.table
        lines.append(f"{indent}  [TABLE rows={len(table.rows)} cols={len(table.columns)}]")
        for r_idx, row in enumerate(table.rows):
            cell_texts = [cell.text for cell in row.cells]
            lines.append(f"{indent}    row{r_idx}: {cell_texts}")
    if shape.is_placeholder:
        lines.append(
            f"{indent}  [PLACEHOLDER idx={shape.placeholder_format.idx} "
            f"type={shape.placeholder_format.type}]"
        )
    if shape.shape_type == 6:  # GROUP
        lines.append(f"{indent}  [GROUP with {len(shape.shapes)} sub-shapes]")
        for sub in shape.shapes:
            lines.extend(describe_shape(sub, indent + "    "))
    return lines


def main(path: str) -> None:
    prs = Presentation(path)
    out_lines: list[str] = []
    out_lines.append(f"FILE: {path}")
    out_lines.append(f"Slide size: {Emu(prs.slide_width).inches:.2f}in x {Emu(prs.slide_height).inches:.2f}in")
    out_lines.append(f"Number of slides: {len(prs.slides.__iter__.__self__._sldIdLst)}")
    out_lines.append(f"Number of slide layouts in slide master: {len(prs.slide_masters[0].slide_layouts)}")

    for s_idx, slide in enumerate(prs.slides, start=1):
        out_lines.append("")
        out_lines.append("=" * 90)
        out_lines.append(f"SLIDE {s_idx}  layout={slide.slide_layout.name!r}")
        out_lines.append("=" * 90)
        for shape in slide.shapes:
            out_lines.extend(describe_shape(shape))

    report = "\n".join(out_lines)
    out_path = Path(__file__).with_name("pptx_dump.txt")
    out_path.write_text(report, encoding="utf-8")
    print(report)
    print(f"\n\n[Written to {out_path}]")


if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else r"c:\Users\JITESH\Downloads\SIH26027_Bug_Hunters_2.pptx"
    main(target)
